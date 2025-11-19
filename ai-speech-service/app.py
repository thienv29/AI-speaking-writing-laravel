from fastapi import FastAPI, UploadFile, File, Form
from fastapi.responses import FileResponse, JSONResponse
from pydantic import BaseModel
from gtts import gTTS
import speech_recognition as sr
from pydub import AudioSegment
import os
import uuid
from fastapi.middleware.cors import CORSMiddleware
from openpyxl import load_workbook

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # hoặc ["http://127.0.0.1:8000"] để an toàn hơn
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Create output directories
os.makedirs("outputs/audio", exist_ok=True)
os.makedirs("outputs/text", exist_ok=True)

class TTSRequest(BaseModel):
    text: str
    lang: str = "en"

@app.post("/tts")
async def tts(req: TTSRequest):
    filename = f"outputs/audio/{uuid.uuid4().hex}.mp3"
    tts = gTTS(req.text, lang=req.lang)
    tts.save(filename)
    print(f"[TTS] Generated: {filename}")
    return FileResponse(filename, media_type="audio/mpeg", filename="output.mp3")

@app.post("/stt")
async def stt(file: UploadFile = File(...)):
    # Validate file
    if not file.filename.lower().endswith((".wav", ".mp3", ".m4a", ".ogg", ".webm")):
        return JSONResponse(status_code=400, content={"error": "Unsupported audio format"})

    # Save uploaded file
    temp_path = f"outputs/audio/{uuid.uuid4().hex}_{file.filename}"
    with open(temp_path, "wb") as f:
        f.write(await file.read())

    # Convert to WAV using pydub
    wav_path = temp_path.rsplit(".", 1)[0] + ".wav"
    try:
        audio = AudioSegment.from_file(temp_path)
        audio.export(wav_path, format="wav")
    except Exception as e:
        return JSONResponse(status_code=500, content={"error": f"Cannot convert audio: {str(e)}"})

    recognizer = sr.Recognizer()
    try:
        with sr.AudioFile(wav_path) as source:
            audio_data = recognizer.record(source)
            text = recognizer.recognize_google(audio_data)
    except sr.UnknownValueError:
        return JSONResponse(status_code=400, content={"error": "Speech could not be recognized"})
    except sr.RequestError:
        return JSONResponse(status_code=503, content={"error": "Speech service unavailable"})
    except Exception as e:
        return JSONResponse(status_code=500, content={"error": str(e)})
    finally:
        for p in [temp_path, wav_path]:
            if os.path.exists(p):
                os.remove(p)

    return JSONResponse(content={"recognized_text": text})

@app.post("/import-excel")
async def import_excel(file: UploadFile = File(...)):
    # Kiểm tra file
    if not file.filename.endswith(".xlsx"):
        return JSONResponse(status_code=400, content={"error": "File must be .xlsx"})
    
    try:
        workbook = load_workbook(file.file)
        sheet = workbook.active

        # Lấy header
        headers = [cell.value for cell in sheet[1]]

        # Parse từng dòng
        data_list = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            obj = {headers[i]: row[i] for i in range(len(headers))}
            data_list.append(obj)

        return {"total": len(data_list), "data": data_list}

    except Exception as e:
        return JSONResponse(status_code=500, content={"error": f"Cannot read Excel file: {str(e)}"})
